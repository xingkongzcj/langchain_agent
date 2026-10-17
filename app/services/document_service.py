"""文档入库服务：解析、切分、向量化。"""

import os
import tempfile
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from fastapi import UploadFile
from langchain_core.documents import Document as LangchainDocument
from langchain_text_splitters import MarkdownHeaderTextSplitter, RecursiveCharacterTextSplitter
from pypdf import PdfReader
from sqlalchemy.orm import Session

from app.core.exceptions import BusinessException
from app.db.chroma import get_vectorstore
from app.models.document import Document as DocumentModel

# 支持的文件扩展名（.doc 为旧版二进制格式，需先另存为 .docx）
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".md", ".txt"}

# 切分参数
CHUNK_SIZE = 500
CHUNK_OVERLAP = 50

# 上传大小限制
MAX_UPLOAD_MB = 20
MAX_UPLOAD_SIZE = MAX_UPLOAD_MB * 1024 * 1024

# Markdown 标题切分规则
MARKDOWN_HEADERS = [
    ("#", "h1"),
    ("##", "h2"),
    ("###", "h3"),
]


def _load_text(file_path: str) -> list[LangchainDocument]:
    """读取纯文本 / Markdown。

    注意：Markdown 保留原始 # 标题语法，以便 MarkdownHeaderTextSplitter 识别章节层级。
    """
    text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
    return [LangchainDocument(page_content=text, metadata={"source": file_path})]


def _load_pdf(file_path: str) -> list[LangchainDocument]:
    """逐页提取 PDF 文本（pypdf，纯 Python，无需 torch/OCR 模型）。"""
    reader = PdfReader(file_path)
    docs: list[LangchainDocument] = []
    for page_index, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(
                LangchainDocument(
                    page_content=text,
                    metadata={"source": file_path, "page": page_index},
                )
            )
    return docs


def _load_docx(file_path: str) -> list[LangchainDocument]:
    """提取 Word（.docx）正文段落与表格文本（python-docx）。"""
    document = DocxDocument(file_path)
    parts: list[str] = [p.text for p in document.paragraphs if p.text.strip()]
    # 表格转成「单元格 | 单元格」的文本行，避免表格内容丢失
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return [LangchainDocument(page_content="\n".join(parts), metadata={"source": file_path})]


def _load_document(file_path: str) -> list[LangchainDocument]:
    """根据扩展名解析文档为 LangChain Document 列表。"""
    ext = Path(file_path).suffix.lower()
    if ext in {".txt", ".md"}:
        return _load_text(file_path)
    if ext == ".pdf":
        return _load_pdf(file_path)
    if ext == ".docx":
        return _load_docx(file_path)
    raise BusinessException(
        code=400,
        message=f"不支持的文件类型: {ext or '未知'}（若为 .doc，请另存为 .docx 后上传）",
    )


def _split_markdown(raw_docs: list[LangchainDocument]) -> list[LangchainDocument]:
    """Markdown 两层切分：先按标题切，再按长度兜底切分。"""
    header_splitter = MarkdownHeaderTextSplitter(
        headers_to_split_on=MARKDOWN_HEADERS,
        strip_headers=False,
    )
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )

    chunks: list[LangchainDocument] = []
    for raw_doc in raw_docs:
        # 第一层：按标题切分，块带 h1/h2/h3 metadata
        header_splits = header_splitter.split_text(raw_doc.page_content)
        # 第二层：对每个章节块按长度兜底切分（保留标题 metadata）
        for header_split in header_splits:
            chunks.extend(text_splitter.split_documents([header_split]))
    return chunks


def _split_generic(raw_docs: list[LangchainDocument]) -> list[LangchainDocument]:
    """通用切分（PDF / Word / TXT）。"""
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
    )
    return text_splitter.split_documents(raw_docs)


def split_documents(
    raw_docs: list[LangchainDocument],
    file_type: str,
) -> list[LangchainDocument]:
    """按文件类型选择切分策略。"""
    if file_type == "md":
        return _split_markdown(raw_docs)
    return _split_generic(raw_docs)


def _build_metadata(
    chunk: LangchainDocument,
    index: int,
    course_id: int,
    document_id: int,
    filename: str,
) -> dict[str, Any]:
    """构建单个 chunk 的 metadata（保留 source 与章节标题）。"""
    meta: dict[str, Any] = {
        "course_id": course_id,
        "document_id": document_id,
        "chunk_index": index,
        "filename": filename,
        "source": filename,
    }
    # 保留章节标题（仅 Markdown 文件有 h1/h2/h3）
    for header_key in ("h1", "h2", "h3"):
        value = chunk.metadata.get(header_key)
        if value:
            meta[header_key] = value
    return meta


def ingest_document(
    db: Session,
    course_id: int,
    file_path: str,
    filename: str,
    file_type: str,
    file_size: int | None = None,
) -> DocumentModel:
    """文档入库：解析 -> 切分 -> 向量化 -> 写库。

    :param db: 数据库会话
    :param course_id: 课程 ID
    :param file_path: 本地文件路径
    :param filename: 原始文件名
    :param file_type: 扩展名（不含点，如 pdf/docx/md/txt）
    :param file_size: 文件大小（字节）
    :return: 入库后的文档记录
    """
    # 1. 创建文档记录（status=processing）
    doc = DocumentModel(
        course_id=course_id,
        filename=filename,
        file_type=file_type,
        file_size=file_size,
        chunk_count=0,
        status="processing",
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    try:
        # 2. 解析文档
        raw_docs = _load_document(file_path)
        if not raw_docs:
            raise BusinessException(code=400, message="文档内容为空，无法入库")

        # 3. 切分文档（Markdown 走两层切分，其余走通用切分）
        chunks = split_documents(raw_docs, file_type)
        if not chunks:
            raise BusinessException(code=400, message="文档切分后无有效内容")

        # 4. 向量化存储到 Chroma
        vectorstore = get_vectorstore(f"course_{course_id}")
        texts = [chunk.page_content for chunk in chunks]
        metadatas = [
            _build_metadata(chunk, index, course_id, doc.id, filename)
            for index, chunk in enumerate(chunks)
        ]
        vectorstore.add_texts(texts=texts, metadatas=metadatas)

        # 5. 更新文档状态为已完成
        doc.chunk_count = len(chunks)
        doc.status = "completed"
        db.commit()
        db.refresh(doc)
        return doc
    except Exception as exc:
        # 解析/向量化失败，标记失败并抛出业务异常
        doc.status = "failed"
        doc.error = str(exc)
        db.commit()
        db.refresh(doc)
        raise BusinessException(code=500, message=f"文档入库失败: {exc}") from exc


async def ingest_upload_file(
    db: Session,
    course_id: int,
    upload_file: UploadFile,
) -> DocumentModel:
    """处理 FastAPI 上传的文件并入库。

    :param db: 数据库会话
    :param course_id: 课程 ID
    :param upload_file: FastAPI 上传的文件对象
    :return: 入库后的文档记录
    """
    filename = upload_file.filename or "unnamed"
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        hint = "（.doc 请另存为 .docx 后上传）" if ext == ".doc" else ""
        raise BusinessException(code=400, message=f"不支持的文件类型: {ext or '未知'}{hint}")

    # 限制单文件大小，避免小内存服务器被大文件上传打爆
    if upload_file.size and upload_file.size > MAX_UPLOAD_SIZE:
        raise BusinessException(code=400, message=f"文件超过 {MAX_UPLOAD_MB}MB 限制")

    content = await upload_file.read()
    if len(content) > MAX_UPLOAD_SIZE:
        raise BusinessException(code=400, message=f"文件超过 {MAX_UPLOAD_MB}MB 限制")

    # 保存到临时文件（Loader 需要文件路径）
    with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp_file:
        tmp_path = tmp_file.name
        tmp_file.write(content)

    try:
        return ingest_document(
            db=db,
            course_id=course_id,
            file_path=tmp_path,
            filename=filename,
            file_type=ext.lstrip("."),
            file_size=len(content),
        )
    finally:
        # 清理临时文件
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def list_documents(db: Session, course_id: int) -> list[DocumentModel]:
    """获取课程下的文档列表。"""
    return (
        db.query(DocumentModel)
        .filter(DocumentModel.course_id == course_id)
        .order_by(DocumentModel.id.desc())
        .all()
    )
